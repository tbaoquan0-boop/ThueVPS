# -*- coding: utf-8 -*-
"""
Bộ tạo tài liệu tiểu luận học thuật: "NGHIÊN CỨU, THIẾT KẾ VÀ HIỆN THỰC HÓA HỆ THỐNG WEBSITE CUNG CẤP DỊCH VỤ MÁY CHỦ ẢO VÀ HẠ TẦNG ĐIỆN TOÁN ĐÁM MÂY VPSSIEUTOC.VN"
Đạt độ dài trên 25 trang A4, định dạng Unicode cỡ chữ 12pt chuẩn đại học Việt Nam.
Tuân thủ tuyệt đối quy tắc: không dùng ngôn từ AI, viết mạch lạc như con người, không dùng gạch ngang đầu dòng, không chấm đầu dòng.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='003366', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), str(edge_data.get('space', 0)))
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(2.0)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Tiểu luận tốt nghiệp chuyên ngành Công nghệ Thông tin")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(9)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun1 = fp.add_run("Hệ thống Web VPSSIEUTOC.VN  |  Trang ")
        frun1.font.name = "Times New Roman"
        frun1.font.size = Pt(9.5)
        frun1.font.color.rgb = RGBColor(120, 120, 120)
        
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
        instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
        fp.add_run()._r.extend([fldChar1, instrText, fldChar2, fldChar3])

def add_p(doc, text, style='body'):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(6)
    
    if style == 'body':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.0)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(30, 30, 30)
    elif style == 'lead':
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.0)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.italic = True
        r.font.color.rgb = RGBColor(40, 40, 40)
    elif style == 'h1':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 51, 102)
    elif style == 'h2':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 60, 120)
    elif style == 'h3':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.italic = True
        r.font.color.rgb = RGBColor(25, 75, 140)
    return p

def add_figure(doc, img_rel_path, caption_text):
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    full_path = os.path.join(base_dir, 'docs', 'img', img_rel_path)
    
    if os.path.exists(full_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        r_img.add_picture(full_path, width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(60, 60, 60)
    else:
        print(f"Warning: Image not found at {full_path}")

print("Helper definitions loaded successfully.")
